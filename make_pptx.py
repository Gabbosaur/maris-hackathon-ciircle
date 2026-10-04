#!/usr/bin/env python3
"""Genera il PowerPoint del pitch MARIS, tarato su 4 minuti.
6 slide, testo essenziale, note con lo script parlato e i tempi.

Le altezze delle caselle sono generose e gli spazi verticali ampi:
in PowerPoint il testo si adatta, ma le caselle non devono
sovrapporsi se qualcuno modifica il contenuto.
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "MARIS-pitch.pptx")

ABYSS = RGBColor(0x05, 0x08, 0x0F)
PANEL = RGBColor(0x0E, 0x17, 0x29)
KELP = RGBColor(0x14, 0xD3, 0xA8)
KELP_L = RGBColor(0x6D, 0xF0, 0xD2)
WHITE = RGBColor(0xF2, 0xF6, 0xFA)
DIM = RGBColor(0xA9, 0xB6, 0xC7)
FAINT = RGBColor(0x7C, 0x8A, 0x9C)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
ORANGE = RGBColor(0xFB, 0x92, 0x3C)

CM = 360000
SLIDE_W = 33.87
SLIDE_H = 19.05

prs = Presentation()
prs.slide_width = Emu(int(SLIDE_W * CM))
prs.slide_height = Emu(int(SLIDE_H * CM))


def cm(v):
    return Emu(int(v * CM))


def blank():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = ABYSS
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def text(slide, x, y, w, h, content, size=18, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, space=8, line=1.3, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    rows = content if isinstance(content, list) else [content]
    for i, row in enumerate(rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = line
        txt, opts = row if isinstance(row, tuple) else (row, {})
        r = p.add_run()
        r.text = txt
        fo = r.font
        fo.name = "Segoe UI"
        fo.size = Pt(opts.get("size", size))
        fo.bold = opts.get("bold", bold)
        fo.color.rgb = opts.get("color", color)
    return tb


def panel(slide, x, y, w, h, fill=PANEL, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cm(x), cm(y), cm(w), cm(h))
    sh.adjustments[0] = 0.05
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def bar(slide, x, y, w=2.0, h=0.14, color=KELP):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cm(x), cm(y), cm(w), cm(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def eyebrow(slide, x, y, label, w=14):
    text(slide, x, y, min(w, SLIDE_W - x - 0.8), 0.7,
         label.upper(), size=12, color=KELP, bold=True, space=0)


def title(slide, label, heading, sub=None):
    bar(slide, 1.8, 1.6)
    eyebrow(slide, 1.8, 2.2, label)
    text(slide, 1.8, 3.0, 30, 2.2, heading, size=38, bold=True,
         color=WHITE, line=1.15, space=0)
    if sub:
        text(slide, 1.8, 5.5, 30, 1.0, sub, size=16, color=DIM, space=0)


def footer(slide, n):
    text(slide, 1.8, 17.9, 24, 0.6,
         "MARIS · Hackathon CIIRCLE · Sfida 2 Robotica sottomarina",
         size=11, color=FAINT, space=0)
    text(slide, 29.0, 17.9, 3.0, 0.6, f"{n}/6", size=11, color=FAINT,
         align=PP_ALIGN.RIGHT, space=0)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


# ---------------------------------------------------------------- 1. titolo
s = blank()
bar(s, 1.8, 4.2, w=2.4, h=0.18)
text(s, 1.8, 4.9, 30, 4.2, "MARIS", size=72, bold=True, color=WHITE, space=0)
text(s, 1.8, 9.3, 30, 1.4,
     "Copilota AI per l'ispezione delle opere marittime",
     size=26, color=KELP_L, space=0)
text(s, 1.8, 11.0, 30, 1.0,
     "Da ore di video a referto tecnico, con memoria storica del degrado",
     size=16, color=DIM, space=0)

panel(s, 1.8, 12.6, 14.4, 3.8)
eyebrow(s, 2.6, 13.2, "Team 7", w=12)
text(s, 2.6, 14.1, 12.8, 2.0,
     ["Lorenzo Torre · Alessandro Melillo · Gabriele Guo"],
     size=15, color=DIM, space=5)

panel(s, 17.6, 12.6, 14.4, 3.8)
eyebrow(s, 18.4, 13.2, "Sfida", w=12)
text(s, 18.4, 14.1, 12.8, 2.0,
     ["Robotica sottomarina e tecnologie subacquee",
      "Salone Nautico di Genova · 3-4 ottobre 2026"],
     size=14, color=DIM, space=5)
footer(s, 1)
notes(s, """[0:00-0:10] Apertura secca.

"Siamo il Team 7. MARIS: un copilota AI per l'ispezione
delle infrastrutture marittime."

Non leggere i nomi uno per uno: sono sulla slide.
Passa subito al problema.""")

# ---------------------------------------------------------------- 2. problema
s = blank()
title(s, "Il problema", "Ogni ispezione nasce e muore in un video")

rows = [
    ("Ore di video riviste a mano",
     "Un tecnico specializzato guarda tutto, fotogramma per fotogramma"),
    ("Valutazione soggettiva",
     "Lo stesso difetto cambia giudizio da operatore a operatore"),
    ("Nessuna memoria storica",
     "Non si sa se una lesione è peggiorata dall'ultima ispezione"),
]
y = 6.6
for head, sub in rows:
    panel(s, 1.8, y, 18.2, 2.9)
    bar(s, 1.8, y, w=0.16, h=2.9, color=ORANGE)
    text(s, 2.8, y + 0.5, 16.6, 1.0, head, size=19, bold=True, color=WHITE, space=0)
    text(s, 2.8, y + 1.6, 16.6, 1.1, sub, size=13, color=FAINT, space=0)
    y += 3.3

panel(s, 21.4, 6.6, 10.6, 9.5, line=KELP)
eyebrow(s, 22.2, 7.2, "In Liguria", w=9)
text(s, 22.2, 8.2, 9.0, 8.8,
     [("Porti di Genova, Savona-Vado, La Spezia", {"size": 15, "color": WHITE}),
      ("Nuova diga foranea di Genova: monitoraggio per decenni",
       {"size": 14, "color": DIM}),
      ("Mareggiata 2018: opere costiere devastate", {"size": 14, "color": DIM}),
      ("Manutenzione reattiva, non predittiva", {"size": 14, "color": AMBER})],
     space=16, line=1.35)
footer(s, 2)
notes(s, """[0:10-0:40] Il problema. Concreto, non generico.

"Le banchine e le dighe liguri si ispezionano con ROV e
sommozzatori. Ogni ispezione produce ore di video che un
tecnico rivede a mano. La valutazione è soggettiva. E
soprattutto: non esiste memoria storica. Nessuno sa se
una lesione è peggiorata rispetto all'anno scorso."

Se hai un dato reale (costo o frequenza ispezioni), dillo qui.
Chiudi con il tocco personale: "sono ligure, questo problema
lo vedo ogni volta che passo in porto".""")

# ---------------------------------------------------------------- 3. architettura
s = blank()
title(s, "Come funziona", "Un motore, quattro asset",
      "L'operatore sceglie il caso d'uso: cambiano prompt di dominio e base normativa")

steps = [
    ("1", "Caso d'uso", "Carica prompt\ne normativa"),
    ("2", "Sito e storico", "Sceglie la baseline\ndi confronto"),
    ("3", "Telemetria", "Quota, temperatura,\ntorbidità dal log"),
    ("4", "Analisi VLM", "Osservazione,\ndiagnosi, azione"),
    ("5", "Referto", "Severità, scadenza,\nfonte citata"),
]
x = 1.8
for n, head, sub in steps:
    panel(s, x, 7.2, 5.6, 5.2)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, cm(x + 0.7), cm(7.8), cm(1.2), cm(1.2))
    circ.fill.solid()
    circ.fill.fore_color.rgb = KELP
    circ.line.fill.background()
    circ.shadow.inherit = False
    tf = circ.text_frame
    tf.text = n
    pr = tf.paragraphs[0]
    pr.alignment = PP_ALIGN.CENTER
    pr.runs[0].font.size = Pt(16)
    pr.runs[0].font.bold = True
    pr.runs[0].font.color.rgb = ABYSS
    pr.runs[0].font.name = "Segoe UI"
    text(s, x + 0.7, 9.4, 4.4, 1.0, head, size=15, bold=True, color=WHITE, space=0)
    text(s, x + 0.7, 10.8, 4.4, 1.8, sub, size=12.5, color=FAINT, space=0, line=1.3)
    x += 6.05

panel(s, 1.8, 13.2, 14.6, 3.4, line=KELP)
eyebrow(s, 2.6, 13.8, "Il differenziatore", w=12)
text(s, 2.6, 14.7, 13.0, 1.6,
     "Ogni raccomandazione è ancorata a una fonte normativa, non inventata dal modello",
     size=15, color=KELP_L, space=0, line=1.35)

panel(s, 17.6, 13.2, 14.4, 3.4)
eyebrow(s, 18.4, 13.8, "Asset coperti", w=12)
text(s, 18.4, 14.7, 12.8, 1.6,
     "Banchine · Dighe frangiflutti · Cavi e condotte · Scafi",
     size=15, color=DIM, space=0, line=1.35)
footer(s, 3)
notes(s, """[0:40-1:00] Architettura. Veloce, non ti perdere qui.

"L'operatore sceglie il tipo di ispezione. Questo carica un
prompt di dominio e una base normativa specifica. Il video
arriva con la telemetria dal log del ROV. Il modello analizza,
e per ogni anomalia produce osservazione, diagnosi e azione."

Il punto da marcare: "non chiediamo a un LLM cosa pensa.
Colleghiamo osservazione visiva e norma tecnica."

Non è detection: quello è risolto e saturo. È comprensione.""")

# ---------------------------------------------------------------- 4. demo
s = blank()
title(s, "Demo", "Ispezione di un cavo sottomarino")

panel(s, 1.8, 6.4, 19.2, 10.0, line=KELP)
text(s, 2.6, 7.2, 17.6, 1.1, "[ VIDEO DELLA DEMO · 60-75 s ]",
     size=19, bold=True, color=KELP, align=PP_ALIGN.CENTER, space=0)
text(s, 2.6, 8.8, 17.6, 7.6,
     [("Inserisci qui il video registrato", {"size": 14, "color": DIM}),
      ("Percorso: caso d'uso, sito e baseline, video, referto",
       {"size": 13, "color": FAINT}),
      ("Da mostrare in sequenza:", {"size": 13.5, "color": WHITE, "bold": True}),
      ("Il prompt che cambia tra banchine e cavi", {"size": 13, "color": DIM}),
      ("Telemetria sincronizzata: quota, temperatura, torbidità",
       {"size": 13, "color": DIM}),
      ("Anomalie che compaiono man mano, non tutte alla fine",
       {"size": 13, "color": DIM}),
      ("Click sull'anomalia: il video torna sul fotogramma",
       {"size": 13, "color": DIM})],
     space=11, line=1.3)

panel(s, 22.4, 6.4, 9.6, 4.8, line=ORANGE)
eyebrow(s, 23.2, 7.0, "Momento chiave", w=8)
text(s, 23.2, 8.0, 8.0, 3.4,
     [("Confronto con l'ispezione precedente", {"size": 14, "color": DIM}),
      ("+6 m di cavo scoperto", {"size": 19, "color": ORANGE, "bold": True})],
     space=10, line=1.3)

panel(s, 22.4, 11.6, 9.6, 4.8)
eyebrow(s, 23.2, 12.2, "Analisi reale", w=8)
text(s, 23.2, 13.2, 8.0, 2.8,
     "Il video dei cavi è stato analizzato con un VLM usando il nostro prompt di dominio",
     size=14, color=DIM, space=0, line=1.35)
footer(s, 4)
notes(s, """[1:00-2:20] LA DEMO. Un terzo del tempo: è ciò che convince.

USA IL VIDEO REGISTRATO, non il live. Con 4 minuti non puoi
rischiare un crash o un caricamento lento.

Commenta mentre scorre, non stare in silenzio:
- "qui seleziona il caso d'uso: cambia il prompt"
- "la telemetria arriva dal log del ROV"
- "le anomalie compaiono man mano"
- "clicco: il video torna sul fotogramma e il riquadro
   circonda l'oggetto"

Poi il momento chiave, rallenta la voce:
"Questa è la parte che conta: il confronto con l'ispezione
precedente. Sei metri di cavo in più scoperti rispetto a
febbraio. Da ispezione isolata a manutenzione predittiva."

ONESTÀ: "i referti in demo sono precaricati, ma l'analisi di
questo video è stata fatta davvero da un VLM con il nostro
prompt." Dillo tu prima che lo chiedano.""")

# ---------------------------------------------------------------- 5. pro e sfide
s = blank()
title(s, "Punti di forza e limiti", "Cosa funziona, cosa resta aperto")

panel(s, 1.8, 6.4, 14.6, 10.0, line=KELP)
eyebrow(s, 2.6, 7.0, "Punti di forza", w=12)
pros = [
    "Nessun addestramento: si compone, non si allena",
    "Zero-shot: descrive anche difetti mai visti",
    "Raccomandazioni ancorate a normativa",
    "Memoria storica: il degrado nel tempo",
    "Un motore, quattro asset: scalabile",
]
y = 8.1
for item in pros:
    bar(s, 2.6, y + 0.32, w=0.26, h=0.26, color=KELP)
    text(s, 3.4, y, 12.2, 1.4, item, size=15, color=WHITE, space=0, line=1.3)
    y += 1.65

panel(s, 17.4, 6.4, 14.6, 10.0, line=AMBER)
eyebrow(s, 18.2, 7.0, "Sfide", w=12)
cons = [
    ("Allucinazioni del modello",
     "Copilota, non giudice: l'ispettore valida sempre"),
    ("Visibilità variabile",
     "Il sistema dichiara la confidenza e quando è bassa"),
    ("Responsabilità tecnica",
     "Prepariamo la decisione, la firma resta all'ingegnere"),
]
y = 8.1
for head, fix in cons:
    text(s, 18.2, y, 13.0, 0.9, head, size=16, bold=True, color=AMBER, space=0)
    text(s, 18.2, y + 1.05, 13.0, 1.6, fix, size=13.5, color=DIM, space=0, line=1.3)
    y += 2.75
footer(s, 5)
notes(s, """[2:20-3:10] Forza e limiti. Qui dimostri maturità.

Sui punti di forza non leggere tutto, scegli tre:
"Non addestriamo niente: usiamo modelli esistenti con un
prompt di dominio. Funziona zero-shot, quindi descrive anche
difetti che non ha mai visto. E ogni raccomandazione cita
una fonte normativa."

Sulle sfide, affrontale tu per primo:
"Un modello linguistico può allucinare. Per questo il sistema
propone e l'ispettore valida: è un copilota, non un giudice.
Quando la visibilità è scarsa il sistema lo dichiara e abbassa
la confidenza. E la responsabilità tecnica resta all'ingegnere:
noi prepariamo la decisione, non la prendiamo."

Dichiarare i limiti prima che te li chiedano vale molto.""")

# ---------------------------------------------------------------- 6. prossimi passi
s = blank()
title(s, "Prossimi passi", "Dove può arrivare")

phases = [
    ("Oggi", "Prototipo funzionante\nReferto da video e storico", KELP, True),
    ("Breve", "VLM in produzione\nNormativa reale indicizzata", KELP_L, False),
    ("Medio", "Dati 3D esistenti\nDegrado in millimetri", DIM, False),
    ("Visione", "Analisi a bordo del ROV\nMonitoraggio continuo", FAINT, False),
]
x = 1.8
for label, body, col, filled in phases:
    panel(s, x, 6.8, 7.2, 5.4, line=col if filled else None)
    text(s, x + 0.8, 7.5, 5.8, 0.9, label.upper(), size=14, bold=True,
         color=col, space=0)
    text(s, x + 0.8, 8.8, 5.8, 3.0, body,
         size=14, color=WHITE if filled else DIM, space=0, line=1.35)
    x += 7.6

panel(s, 1.8, 13.0, 14.6, 3.6)
eyebrow(s, 2.6, 13.6, "Altre applicazioni", w=12)
text(s, 2.6, 14.6, 13.0, 1.8,
     "Scafi e refit · Impianti offshore · Acquacoltura · Monitoraggio ambientale",
     size=15, color=DIM, space=0, line=1.4)

panel(s, 17.4, 13.0, 14.6, 3.6, line=KELP)
eyebrow(s, 18.2, 13.6, "Il passo successivo", w=12)
text(s, 18.2, 14.6, 13.0, 1.8,
     "Pilota con un'Autorità di Sistema Portuale ligure, verso la Call for solution CIIRCLE",
     size=15, color=KELP_L, space=0, line=1.4)
footer(s, 6)
notes(s, """[3:10-4:00] Chiusura. Guarda la giuria, non la slide.

"Oggi abbiamo un prototipo che funziona. Nel breve colleghiamo
un VLM in produzione e indicizziamo normativa reale. Nel medio
estendiamo ai dati 3D che i porti già producono, per quantificare
il degrado in millimetri. La visione è portare l'analisi a bordo
del ROV: sott'acqua la banda non c'è, quindi non trasmetti il
video, trasmetti la comprensione."

Chiudi forte:
"Il passo immediato è un pilota con un'Autorità Portuale ligure.
È la base naturale per la Call for solution del progetto CIIRCLE."

Ultima frase, lenta:
"Oggi ogni ispezione nasce e muore in un video. Noi le diamo
memoria."

Poi silenzio. Aspetta le domande.""")

prs.save(OUT)
print(f"Creato: {OUT}")
