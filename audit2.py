"""Stima l'altezza reale del testo e trova dove invade la casella sottostante."""

from pptx import Presentation

prs = Presentation("MARIS-pitch.pptx")
CM = 360000

# larghezza media di un carattere in cm, per punto di font (Segoe UI)
CHAR_W_PER_PT = 0.0185
# altezza riga in cm per punto di font, con line spacing
LINE_H_PER_PT = 0.0415


def est_height(sh):
    """Altezza stimata del contenuto, considerando il wrap."""
    w_cm = sh.width / CM
    total = 0.0
    for p in sh.text_frame.paragraphs:
        txt = "".join(r.text for r in p.runs)
        size = 18
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        # righe esplicite (newline) + wrap
        lines = 0
        for seg in txt.split("\n"):
            char_w = size * CHAR_W_PER_PT
            per_line = max(1, int(w_cm / char_w))
            lines += max(1, -(-len(seg) // per_line))
        total += lines * size * LINE_H_PER_PT * ls
        if p.space_after:
            total += p.space_after.pt * 0.0353
    return total


print("=" * 74)
print("CASELLE IL CUI TESTO ECCEDE L'ALTEZZA DICHIARATA")
print("=" * 74)

problems = []
for i, s in enumerate(prs.slides, 1):
    items = []
    for sh in s.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        if sh.left is None:
            continue
        h_dec = sh.height / CM
        h_est = est_height(sh)
        items.append(
            {
                "t": sh.text_frame.text.strip().replace("\n", " ⏎ "),
                "l": sh.left / CM,
                "r": (sh.left + sh.width) / CM,
                "top": sh.top / CM,
                "b_dec": sh.top / CM + h_dec,
                "b_est": sh.top / CM + h_est,
                "h_dec": h_dec,
                "h_est": h_est,
            }
        )
        if h_est > h_dec + 0.25:
            print(f"\nslide {i}: dichiarata {h_dec:.1f}cm, stimata {h_est:.1f}cm (+{h_est - h_dec:.1f})")
            print(f"  '{sh.text_frame.text.strip()[:70]}'")
            problems.append((i, sh.text_frame.text.strip()[:40]))

    # ora cerco invasioni reali usando l'altezza stimata
    for a in items:
        for b in items:
            if a is b:
                continue
            # b sta sotto a, nella stessa colonna?
            same_col = min(a["r"], b["r"]) - max(a["l"], b["l"]) > 1.0
            if not same_col or b["top"] < a["top"]:
                continue
            invade = a["b_est"] - b["top"]
            if invade > 0.2:
                print(f"\nslide {i}: INVASIONE di {invade:.1f}cm")
                print(f"  sopra: '{a['t'][:50]}' finisce a {a['b_est']:.1f}")
                print(f"  sotto: '{b['t'][:50]}' inizia a {b['top']:.1f}")
                problems.append((i, "invasione"))

print(f"\n{'=' * 74}")
print(f"Totale problemi: {len(problems)}")
