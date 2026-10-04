"""Controlla sovrapposizioni tra testi e problemi ortografici nel pptx."""

import re

from pptx import Presentation

prs = Presentation("MARIS-pitch.pptx")
CM = 360000
SW = prs.slide_width / CM
SH = prs.slide_height / CM

# parole che devono avere l'accento
ACCENTI = {
    "severita": "severità",
    "torbidita": "torbidità",
    "normativita": "normatività",
    "puo": "può",
    "e'": "è",
    "scegle": "sceglie",
    "Autorita": "Autorità",
    "autorita": "autorità",
    "responsabilita": "responsabilità",
    "visibilita": "visibilità",
    "attivita": "attività",
    "qualita": "qualità",
    "citta": "città",
    "perche": "perché",
    "piu": "più",
    "cio": "ciò",
    "gia": "già",
    "cosi": "così",
    "unita": "unità",
    "proprieta": "proprietà",
    "meta": "metà",
    "sara": "sarà",
    "verra": "verrà",
    "cambiera": "cambierà",
}


def boxes(slide):
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if not txt:
            continue
        if sh.left is None or sh.width is None:
            continue
        out.append(
            {
                "t": txt,
                "l": sh.left / CM,
                "r": (sh.left + sh.width) / CM,
                "top": sh.top / CM,
                "b": (sh.top + sh.height) / CM,
            }
        )
    return out


def overlap(a, b):
    ox = min(a["r"], b["r"]) - max(a["l"], b["l"])
    oy = min(a["b"], b["b"]) - max(a["top"], b["top"])
    if ox <= 0.05 or oy <= 0.05:
        return None
    return ox * oy, ox, oy


print("=" * 72)
print("SOVRAPPOSIZIONI TRA CASELLE DI TESTO")
print("=" * 72)
found = 0
for i, s in enumerate(prs.slides, 1):
    bs = boxes(s)
    for x in range(len(bs)):
        for y in range(x + 1, len(bs)):
            ov = overlap(bs[x], bs[y])
            if ov and ov[0] > 0.3:  # oltre 0.3 cm2
                found += 1
                print(f"\nslide {i}: area {ov[0]:.1f} cm2 (x {ov[1]:.1f} · y {ov[2]:.1f})")
                print(f"  A: '{bs[x]['t'][:52]}'")
                print(f"     L{bs[x]['l']:.1f} R{bs[x]['r']:.1f} T{bs[x]['top']:.1f} B{bs[x]['b']:.1f}")
                print(f"  B: '{bs[y]['t'][:52]}'")
                print(f"     L{bs[y]['l']:.1f} R{bs[y]['r']:.1f} T{bs[y]['top']:.1f} B{bs[y]['b']:.1f}")
if not found:
    print("\nNessuna sovrapposizione rilevante")

print("\n" + "=" * 72)
print("ORTOGRAFIA E ACCENTI")
print("=" * 72)
issues = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                t = r.text
                for bad, good in ACCENTI.items():
                    if re.search(r"\b" + re.escape(bad) + r"\b", t, re.IGNORECASE):
                        issues += 1
                        print(f"  slide {i}: '{bad}' -> '{good}'  in: {t[:60]}")
    # anche nelle note
    nt = s.notes_slide.notes_text_frame.text
    for bad, good in ACCENTI.items():
        if re.search(r"\b" + re.escape(bad) + r"\b", nt, re.IGNORECASE):
            issues += 1
            print(f"  slide {i} NOTE: '{bad}' -> '{good}'")

if not issues:
    print("\nNessun problema ortografico")
print(f"\nTotale: {found} sovrapposizioni, {issues} problemi ortografici")
